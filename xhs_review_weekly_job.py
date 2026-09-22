# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import base64
from datetime import datetime
import hashlib
import hmac
import json
import mimetypes
from pathlib import Path
import re
import time
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
import uuid
import zipfile

from encoding_guard import enable_utf8_stdio
from xhs_review_image_downloader import (
    DEFAULT_PORT,
    connect_review_page,
    download_file,
    fetch_detail,
    fetch_review_rows,
    safe_name,
)

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


enable_utf8_stdio()


BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "config" / "feishu_webhook.json"
STATE_FILE = BASE_DIR / "data" / "xhs_review_images" / "weekly_state.json"
DEFAULT_OUT_ROOT = BASE_DIR / "data" / "xhs_review_images" / "weekly"
REVIEW_SHEET_HEADERS = [
    "\u56fe\u7247",
    "\u6635\u79f0",
    "\u5546\u54c1\u540d\u79f0",
    "\u5546\u54c1ID",
    "\u8bc4\u4ef7ID",
    "\u56fe\u7247\u5e8f\u53f7",
    "\u8bc4\u8bba\u5185\u5bb9",
    "\u8bc4\u8bba\u65f6\u95f4",
    "\u539f\u56fe\u94fe\u63a5",
    "\u672c\u5730\u6587\u4ef6",
    "\u72b6\u6001",
]
DEFAULT_SHEET_NAME = "\u8bc4\u8bba\u7d20\u6750"
ANONYMOUS_USER = "\u533f\u540d\u7528\u6237"
PRODUCT_FALLBACK = "\u5546\u54c1"
IMAGE_EMBED_FAILED = "\u56fe\u7247\u5d4c\u5165\u5931\u8d25"


def load_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def save_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_feishu_config() -> dict[str, Any]:
    config = load_json(CONFIG_FILE, {})
    cleaned: dict[str, Any] = {}
    for key, value in config.items():
        if isinstance(value, dict):
            cleaned[str(key)] = value
        else:
            text = str(value).strip()
            if text:
                cleaned[str(key)] = text
    return cleaned


def feishu_sign(timestamp: str, secret: str) -> str:
    string_to_sign = f"{timestamp}\n{secret}".encode("utf-8")
    digest = hmac.new(string_to_sign, b"", digestmod=hashlib.sha256).digest()
    return base64.b64encode(digest).decode("utf-8")


def send_webhook_text(text: str, config: dict[str, str]) -> bool:
    webhook = config.get("webhook", "")
    if not webhook:
        return False
    message: dict[str, Any] = {"msg_type": "text", "content": {"text": text}}
    secret = config.get("secret", "")
    if secret:
        timestamp = str(int(time.time()))
        message["timestamp"] = timestamp
        message["sign"] = feishu_sign(timestamp, secret)
    body = json.dumps(message, ensure_ascii=False).encode("utf-8")
    request = Request(
        webhook,
        data=body,
        headers={"Content-Type": "application/json; charset=utf-8"},
        method="POST",
    )
    with urlopen(request, timeout=20) as response:
        response.read()
    return True


def feishu_json_request(url: str, payload: dict[str, Any], token: str | None = None) -> dict[str, Any]:
    headers = {"Content-Type": "application/json; charset=utf-8"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    with urlopen(request, timeout=30) as response:
        data = json.loads(response.read().decode("utf-8", errors="ignore") or "{}")
    if data.get("code", 0) != 0:
        raise RuntimeError(f"Feishu API failed: {data}")
    return data


def feishu_get_request(url: str, token: str | None = None) -> dict[str, Any]:
    headers = {}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = Request(url, headers=headers, method="GET")
    with urlopen(request, timeout=30) as response:
        data = json.loads(response.read().decode("utf-8", errors="ignore") or "{}")
    if data.get("code", 0) != 0:
        raise RuntimeError(f"Feishu API failed: {data}")
    return data


def get_tenant_access_token(config: dict[str, str]) -> str | None:
    app_id = config.get("app_id", "")
    app_secret = config.get("app_secret", "")
    if not app_id or not app_secret:
        return None
    data = feishu_json_request(
        "https://open.feishu.cn/open-apis/auth/v3/tenant_access_token/internal",
        {"app_id": app_id, "app_secret": app_secret},
    )
    return str(data.get("tenant_access_token") or "")


def multipart_request(url: str, fields: dict[str, str], files: dict[str, Path], token: str) -> dict[str, Any]:
    boundary = "----codex" + uuid.uuid4().hex
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.extend(
            [
                f"--{boundary}\r\n".encode(),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                value.encode("utf-8"),
                b"\r\n",
            ]
        )
    for name, path in files.items():
        content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        chunks.extend(
            [
                f"--{boundary}\r\n".encode(),
                f'Content-Disposition: form-data; name="{name}"; filename="{path.name}"\r\n'.encode(),
                f"Content-Type: {content_type}\r\n\r\n".encode(),
                path.read_bytes(),
                b"\r\n",
            ]
        )
    chunks.append(f"--{boundary}--\r\n".encode())
    request = Request(
        url,
        data=b"".join(chunks),
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
        method="POST",
    )
    with urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8", errors="ignore") or "{}")
    if data.get("code", 0) != 0:
        raise RuntimeError(f"Feishu upload failed: {data}")
    return data


def upload_feishu_image(path: Path, token: str) -> str:
    data = multipart_request(
        "https://open.feishu.cn/open-apis/im/v1/images",
        {"image_type": "message"},
        {"image": path},
        token,
    )
    return str((data.get("data") or {}).get("image_key") or "")


def send_feishu_image(image_key: str, chat_id: str, token: str) -> None:
    feishu_json_request(
        "https://open.feishu.cn/open-apis/im/v1/messages?receive_id_type=chat_id",
        {
            "receive_id": chat_id,
            "msg_type": "image",
            "content": json.dumps({"image_key": image_key}, ensure_ascii=False),
        },
        token,
    )


def upload_feishu_file(path: Path, token: str) -> str:
    data = multipart_request(
        "https://open.feishu.cn/open-apis/im/v1/files",
        {"file_type": "stream", "file_name": path.name},
        {"file": path},
        token,
    )
    return str((data.get("data") or {}).get("file_key") or "")


def send_feishu_file(file_key: str, chat_id: str, token: str) -> None:
    feishu_json_request(
        "https://open.feishu.cn/open-apis/im/v1/messages?receive_id_type=chat_id",
        {
            "receive_id": chat_id,
            "msg_type": "file",
            "content": json.dumps({"file_key": file_key}, ensure_ascii=False),
        },
        token,
    )


def send_feishu_text(chat_id: str, text: str, token: str) -> None:
    feishu_json_request(
        "https://open.feishu.cn/open-apis/im/v1/messages?receive_id_type=chat_id",
        {
            "receive_id": chat_id,
            "msg_type": "text",
            "content": json.dumps({"text": text}, ensure_ascii=False),
        },
        token,
    )


def parse_sheet_url(url: str) -> tuple[str, str]:
    token_match = re.search(r"/sheets/([^/?#]+)", url)
    sheet_match = re.search(r"[?&#]sheet=([^&#]+)", url)
    token = token_match.group(1) if token_match else ""
    sheet_id = sheet_match.group(1) if sheet_match else ""
    return token, sheet_id


def parse_wiki_url(url: str) -> str:
    match = re.search(r"/wiki/([^/?#]+)", url)
    return match.group(1) if match else ""


def first_sheet_id(spreadsheet_token: str, token: str) -> str:
    data = feishu_get_request(
        f"https://open.feishu.cn/open-apis/sheets/v3/spreadsheets/{spreadsheet_token}/sheets/query",
        token,
    )
    sheets = (data.get("data") or {}).get("sheets") or []
    if not sheets:
        return ""
    first = sheets[0]
    return str(first.get("sheet_id") or first.get("sheetId") or "")


def resolve_wiki_sheet(url: str, token: str) -> tuple[str, str]:
    wiki_token = parse_wiki_url(url)
    if not wiki_token:
        return "", ""
    data = feishu_get_request(
        "https://open.feishu.cn/open-apis/wiki/v2/spaces/get_node?token=" + wiki_token,
        token,
    )
    node = (data.get("data") or {}).get("node") or data.get("data") or {}
    obj_token = str(node.get("obj_token") or node.get("objToken") or "")
    obj_type = str(node.get("obj_type") or node.get("objType") or "").lower()
    if obj_type and "sheet" not in obj_type:
        raise RuntimeError(f"Wiki node is not a spreadsheet: obj_type={obj_type}")
    sheet_id = first_sheet_id(obj_token, token) if obj_token else ""
    return obj_token, sheet_id


def sheet_config(config: dict[str, Any]) -> tuple[str, str]:
    token = str(config.get("feishu_sheet_token") or "").strip()
    sheet_id = str(config.get("feishu_sheet_id") or "").strip()
    url = str(config.get("feishu_sheet_url") or "").strip()
    if url and (not token or not sheet_id):
        parsed_token, parsed_sheet_id = parse_sheet_url(url)
        token = token or parsed_token
        sheet_id = sheet_id or parsed_sheet_id
    return token, sheet_id


def append_sheet_values(spreadsheet_token: str, sheet_id: str, rows: list[list[Any]], token: str) -> str:
    data = feishu_json_request(
        f"https://open.feishu.cn/open-apis/sheets/v2/spreadsheets/{spreadsheet_token}/values_append",
        {"valueRange": {"range": f"{sheet_id}!A:K", "values": rows}},
        token,
    )
    payload = data.get("data") or {}
    updates = payload.get("updates") or payload.get("valueRange") or {}
    return str(updates.get("updatedRange") or updates.get("range") or "")


def range_start_row(updated_range: str, fallback: int = 2) -> int:
    match = re.search(r"![A-Z]+(\d+)", updated_range or "")
    if not match:
        return fallback
    return int(match.group(1))


def write_sheet_image(spreadsheet_token: str, sheet_id: str, row: int, image_path: Path, token: str) -> None:
    feishu_json_request(
        f"https://open.feishu.cn/open-apis/sheets/v2/spreadsheets/{spreadsheet_token}/values_image",
        {
            "range": f"{sheet_id}!A{row}:A{row}",
            "image": list(image_path.read_bytes()),
            "name": image_path.name,
        },
        token,
    )


def append_items_to_online_sheet(items: list[dict[str, Any]], config: dict[str, Any], token: str) -> int:
    spreadsheet_token, sheet_id = sheet_config(config)
    sheet_url = str(config.get("feishu_sheet_url") or "").strip()
    if sheet_url and (not spreadsheet_token or not sheet_id):
        wiki_token, wiki_sheet_id = resolve_wiki_sheet(sheet_url, token)
        spreadsheet_token = spreadsheet_token or wiki_token
        sheet_id = sheet_id or wiki_sheet_id
    if not spreadsheet_token or not sheet_id:
        return 0

    rows: list[list[Any]] = []
    ok_items: list[dict[str, Any]] = []
    for item in items:
        path = Path(str(item.get("file") or ""))
        if not path.exists() or str(item.get("status")) != "ok":
            continue
        ok_items.append(item)
        rows.append(
            [
                "",
                item.get("nickname") or "",
                item.get("product_name") or "",
                item.get("item_id") or "",
                item.get("review_id") or "",
                item.get("image_no") or "",
                item.get("review_text") or "",
                item.get("create_time") or "",
                item.get("image_url") or "",
                item.get("file") or "",
                item.get("status") or "",
            ]
        )

    if not rows:
        return 0

    updated_range = append_sheet_values(spreadsheet_token, sheet_id, rows, token)
    first_row = range_start_row(updated_range, fallback=2)
    for offset, item in enumerate(ok_items):
        write_sheet_image(
            spreadsheet_token,
            sheet_id,
            first_row + offset,
            Path(str(item.get("file") or "")),
            token,
        )
        time.sleep(0.2)
    return len(ok_items)


def upload_bitable_attachment(path: Path, app_token: str, token: str) -> str:
    data = multipart_request(
        "https://open.feishu.cn/open-apis/drive/v1/medias/upload_all",
        {
            "file_name": path.name,
            "parent_type": "bitable_image",
            "parent_node": app_token,
            "size": str(path.stat().st_size),
        },
        {"file": path},
        token,
    )
    return str((data.get("data") or {}).get("file_token") or "")


def create_bitable_records(app_token: str, table_id: str, records: list[dict[str, Any]], token: str) -> int:
    if not records:
        return 0
    total = 0
    for start in range(0, len(records), 100):
        batch = records[start : start + 100]
        data = feishu_json_request(
            f"https://open.feishu.cn/open-apis/bitable/v1/apps/{app_token}/tables/{table_id}/records/batch_create",
            {"records": batch},
            token,
        )
        total += len((data.get("data") or {}).get("records") or batch)
    return total


def bitable_field_map(config: dict[str, Any]) -> dict[str, str]:
    fields = config.get("bitable_fields")
    if not isinstance(fields, dict):
        fields = {}
    defaults = {
        "image": "图片",
        "nickname": "昵称",
        "product_name": "商品名称",
        "item_id": "商品ID",
        "review_id": "评价ID",
        "image_no": "图片序号",
        "review_text": "评论内容",
        "create_time": "评论时间",
        "crawl_time": "抓取时间",
        "image_url": "原图链接",
        "local_file": "本地文件",
    }
    return {key: str(fields.get(key) or value) for key, value in defaults.items()}


def upload_items_to_bitable(
    items: list[dict[str, Any]],
    config: dict[str, Any],
    token: str,
    run_at: str,
) -> int:
    app_token = str(config.get("bitable_app_token") or "").strip()
    table_id = str(config.get("bitable_table_id") or "").strip()
    if not app_token or not table_id:
        return 0

    field = bitable_field_map(config)
    records: list[dict[str, Any]] = []
    for item in items:
        path = Path(str(item.get("file") or ""))
        if not path.exists():
            continue
        file_token = upload_bitable_attachment(path, app_token, token)
        if not file_token:
            continue
        records.append(
            {
                "fields": {
                    field["image"]: [{"file_token": file_token}],
                    field["nickname"]: str(item.get("nickname") or ""),
                    field["product_name"]: str(item.get("product_name") or ""),
                    field["item_id"]: str(item.get("item_id") or ""),
                    field["review_id"]: str(item.get("review_id") or ""),
                    field["image_no"]: int(item.get("image_no") or 0),
                    field["review_text"]: str(item.get("review_text") or ""),
                    field["create_time"]: str(item.get("create_time") or ""),
                    field["crawl_time"]: run_at,
                    field["image_url"]: str(item.get("image_url") or ""),
                    field["local_file"]: str(item.get("file") or ""),
                }
            }
        )
    return create_bitable_records(app_token, table_id, records, token)


def review_image_key(review_id: str, image_no: int, link: str) -> str:
    digest = hashlib.sha1(link.encode("utf-8", errors="ignore")).hexdigest()[:12]
    return f"{review_id}:{image_no}:{digest}"


def make_zip(out_dir: Path) -> Path:
    zip_path = out_dir.with_suffix(".zip")
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in out_dir.rglob("*"):
            if path.is_file():
                zf.write(path, path.relative_to(out_dir.parent))
    return zip_path


def make_review_workbook(out_dir: Path, manifest: list[dict[str, Any]]) -> Path:
    workbook_path = out_dir.with_suffix(".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "评论素材"
    headers = [
        "图片",
        "昵称",
        "商品名称",
        "商品ID",
        "评价ID",
        "图片序号",
        "评论内容",
        "评论时间",
        "原图链接",
        "本地文件",
        "状态",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [18, 18, 34, 18, 24, 10, 42, 18, 42, 48, 14]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    row_index = 2
    for item in manifest:
        path = Path(str(item.get("file") or ""))
        ws.append(
            [
                "",
                item.get("nickname") or "",
                item.get("product_name") or "",
                item.get("item_id") or "",
                item.get("review_id") or "",
                item.get("image_no") or "",
                item.get("review_text") or "",
                item.get("create_time") or "",
                item.get("image_url") or "",
                item.get("file") or "",
                item.get("status") or "",
            ]
        )
        ws.row_dimensions[row_index].height = 92
        for col in range(1, len(headers) + 1):
            ws.cell(row_index, col).alignment = Alignment(vertical="center", wrap_text=True)
        if path.exists() and str(item.get("status")) == "ok":
            try:
                image = ExcelImage(str(path))
                image.width = 110
                image.height = 110
                ws.add_image(image, f"A{row_index}")
            except Exception:
                ws.cell(row_index, 1).value = "图片嵌入失败"
        row_index += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(workbook_path)
    return workbook_path


def make_review_workbook(
    out_dir: Path,
    manifest: list[dict[str, Any]],
    workbook_path: Path | None = None,
    sheet_name: str = "",
) -> Path:
    workbook_path = workbook_path or out_dir.with_suffix(".xlsx")
    headers = REVIEW_SHEET_HEADERS

    if workbook_path.exists():
        wb = load_workbook(workbook_path)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif sheet_name:
            ws = wb.create_sheet(sheet_name)
        else:
            ws = wb.active
    else:
        workbook_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name or DEFAULT_SHEET_NAME

    if ws.max_row == 1 and all(cell.value is None for cell in ws[1]):
        ws.append(headers)
    elif ws.max_row < 1:
        ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        if cell.value is None:
            continue
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [18, 18, 34, 18, 24, 10, 42, 18, 42, 48, 14]
    for index, width in enumerate(widths, 1):
        current_width = ws.column_dimensions[get_column_letter(index)].width or 0
        if current_width < width:
            ws.column_dimensions[get_column_letter(index)].width = width

    row_index = max(ws.max_row + 1, 2)
    for item in manifest:
        path = Path(str(item.get("file") or ""))
        ws.append(
            [
                "",
                item.get("nickname") or "",
                item.get("product_name") or "",
                item.get("item_id") or "",
                item.get("review_id") or "",
                item.get("image_no") or "",
                item.get("review_text") or "",
                item.get("create_time") or "",
                item.get("image_url") or "",
                item.get("file") or "",
                item.get("status") or "",
            ]
        )
        ws.row_dimensions[row_index].height = 92
        for col in range(1, len(headers) + 1):
            ws.cell(row_index, col).alignment = Alignment(vertical="center", wrap_text=True)
        if path.exists() and str(item.get("status")) == "ok":
            try:
                image = ExcelImage(str(path))
                image.width = 110
                image.height = 110
                ws.add_image(image, f"A{row_index}")
            except Exception:
                ws.cell(row_index, 1).value = IMAGE_EMBED_FAILED
        row_index += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(workbook_path)
    return workbook_path


def collect_new_review_images(port: int, out_dir: Path, baseline: bool = False) -> dict[str, Any]:
    state = load_json(STATE_FILE, {"seen": {}})
    seen: dict[str, Any] = state.setdefault("seen", {})
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, Any]] = []
    page = connect_review_page(port)
    try:
        rows = fetch_review_rows(page)
        for idx, row in enumerate(rows, 1):
            review = row.get("review_data") or {}
            sku = row.get("sku_info") or {}
            review_id = str(review.get("review_id") or "")
            images = ((review.get("content") or {}).get("images") or [])
            if not review_id or not images:
                continue
            detail = fetch_detail(page, review_id, review.get("review_type"))
            user_info = detail.get("user_info") or {}
            nickname = user_info.get("name") or (ANONYMOUS_USER if review.get("anonymous") else "unknown")
            nick_dir = out_dir / safe_name(nickname, "unknown")
            for img_no, img in enumerate(images, 1):
                link = img.get("link") if isinstance(img, dict) else str(img)
                if not link:
                    continue
                key = review_image_key(review_id, img_no, link)
                if key in seen:
                    continue
                seen[key] = {
                    "first_seen_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "review_id": review_id,
                    "image_no": img_no,
                    "image_url": link,
                    "nickname": nickname,
                    "product_name": sku.get("name") or "",
                    "item_id": sku.get("item_id") or "",
                    "create_time": review.get("create_time") or "",
                }
                if baseline:
                    continue
                nick_dir.mkdir(parents=True, exist_ok=True)
                base = f"{idx:04d}_{safe_name(sku.get('name'), PRODUCT_FALLBACK)}_{review_id}_{img_no}"
                try:
                    final_path = download_file(link, nick_dir / base)
                    status = "ok"
                except Exception as exc:
                    final_path = ""
                    status = f"failed: {exc}"
                manifest.append(
                    {
                        "nickname": nickname,
                        "review_id": review_id,
                        "image_no": img_no,
                        "status": status,
                        "file": str(final_path),
                        "image_url": link,
                        "product_name": sku.get("name") or "",
                        "item_id": sku.get("item_id") or "",
                        "sku_id": sku.get("sku_id") or "",
                        "order_id": sku.get("order_id") or "",
                        "review_text": (review.get("content") or {}).get("text") or "",
                        "create_time": review.get("create_time") or "",
                    }
                )
            print(f"已检查 {idx}/{len(rows)}：{review_id}", flush=True)
            time.sleep(0.05)
    finally:
        page.close()

    state["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    save_json(STATE_FILE, state)
    return {"manifest": manifest, "seen_count": len(seen)}


def write_manifest(out_dir: Path, manifest: list[dict[str, Any]]) -> Path:
    path = out_dir / "下载清单.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def notify_feishu(summary: dict[str, Any], config: dict[str, Any], send_images: bool, max_images: int) -> None:
    manifest = summary["manifest"]
    ok_items = [item for item in manifest if item.get("status") == "ok" and item.get("file")]
    failed_items = [item for item in manifest if item.get("status") != "ok"]
    workbook_path = summary.get("workbook_path")
    zip_path = summary.get("zip_path")
    token = get_tenant_access_token(config)
    chat_id = str(config.get("chat_id") or "").strip()
    safe_text = "\n".join(
        [
            "\u5c0f\u7ea2\u4e66\u5e26\u56fe\u8bc4\u4ef7\u7d20\u6750\u6293\u53d6\u5b8c\u6210",
            f"\u6293\u53d6\u65f6\u95f4\uff1a{summary['run_at']}",
            f"\u65b0\u589e\u56fe\u7247\uff1a{len(ok_items)} \u5f20",
            f"\u5931\u8d25\u56fe\u7247\uff1a{len(failed_items)} \u5f20",
            f"\u672c\u5730\u76ee\u5f55\uff1a{summary['out_dir']}",
        ]
    )
    if not send_images or not token or not chat_id:
        print(safe_text)
        return

    if config.get("feishu_sheet_url") or config.get("feishu_sheet_token"):
        inserted = append_items_to_online_sheet(ok_items, config, token)
        send_feishu_text(
            chat_id,
            f"\u5c0f\u7ea2\u4e66\u65b0\u589e\u8bc4\u8bba\u56fe\u7247\u5df2\u5199\u5165\u98de\u4e66\u5728\u7ebf\u8868\u683c\uff1a{inserted} \u6761\u3002",
            token,
        )
        return

    if workbook_path:
        file_key = upload_feishu_file(Path(workbook_path), token)
        if file_key:
            send_feishu_file(file_key, chat_id, token)
        send_feishu_text(chat_id, safe_text, token)
        return

    if zip_path:
        file_key = upload_feishu_file(Path(zip_path), token)
        if file_key:
            send_feishu_file(file_key, chat_id, token)
        send_feishu_text(chat_id, safe_text, token)
        return

    send_feishu_text(chat_id, safe_text, token)
    return
    text = "\n".join(
        [
            "小红书带图评价素材抓取完成",
            f"抓取时间：{summary['run_at']}",
            f"新增图片：{len(ok_items)} 张",
            f"失败图片：{len(failed_items)} 张",
            f"本地目录：{summary['out_dir']}",
            f"打包文件：{zip_path or '-'}",
        ]
    )
    sent_text = send_webhook_text(text, config)

    token = get_tenant_access_token(config)
    chat_id = config.get("chat_id", "")
    if not send_images or not token or not chat_id:
        if not sent_text:
            print(text)
        return

    if config.get("feishu_sheet_url") or config.get("feishu_sheet_token"):
        inserted = append_items_to_online_sheet(ok_items, config, token)
        send_webhook_text(f"小红书新增评论图片已写入飞书在线表格：{inserted} 条。", config)
        return

    if workbook_path:
        file_key = upload_feishu_file(Path(workbook_path), token)
        if file_key:
            send_feishu_file(file_key, chat_id, token)
        return

    if zip_path:
        file_key = upload_feishu_file(Path(zip_path), token)
        if file_key:
            send_feishu_file(file_key, chat_id, token)
        return

    if ok_items:
        send_webhook_text("本次有新增图片，但压缩包未生成成功，请查看本地目录。", config)
    return

    for item in ok_items[:max_images]:
        image_key = upload_feishu_image(Path(item["file"]), token)
        if image_key:
            send_feishu_image(image_key, chat_id, token)
        time.sleep(0.2)

    if len(ok_items) > max_images:
        send_webhook_text(f"本次新增 {len(ok_items)} 张图片，已发送前 {max_images} 张，其余请查看打包文件。", config)


def main() -> int:
    parser = argparse.ArgumentParser(description="每周增量抓取小红书带图评价素材并通知飞书")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    parser.add_argument("--out-root", default=str(DEFAULT_OUT_ROOT))
    parser.add_argument("--baseline", action="store_true", help="只记录当前已有图片为已处理，不下载也不发送")
    parser.add_argument("--no-feishu", action="store_true")
    parser.add_argument("--send-images", action="store_true", help="需要 feishu_webhook.json 配置 app_id/app_secret/chat_id")
    parser.add_argument("--max-images", type=int, default=20)
    args = parser.parse_args()

    run_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = Path(args.out_root) / stamp
    result = collect_new_review_images(args.port, out_dir, baseline=args.baseline)
    manifest = result["manifest"]

    workbook_path = None
    zip_path = None
    if manifest:
        write_manifest(out_dir, manifest)
        config = load_feishu_config()
        target_workbook = str(config.get("review_workbook_path") or "").strip()
        target_sheet = str(config.get("review_workbook_sheet") or "").strip()
        workbook_path = make_review_workbook(
            out_dir,
            manifest,
            Path(target_workbook) if target_workbook else None,
            target_sheet,
        )
        zip_path = make_zip(out_dir)

    summary = {
        "run_at": run_at,
        "out_dir": str(out_dir),
        "workbook_path": str(workbook_path) if workbook_path else "",
        "zip_path": str(zip_path) if zip_path else "",
        "manifest": manifest,
        "seen_count": result["seen_count"],
    }
    print(json.dumps(summary, ensure_ascii=False))

    if not args.baseline and not args.no_feishu:
        try:
            notify_feishu(summary, load_feishu_config(), args.send_images, args.max_images)
        except (HTTPError, URLError, TimeoutError, RuntimeError) as exc:
            print(f"飞书通知失败：{exc}")
            return 2

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
