from pathlib import Path
from openpyxl import load_workbook
import pandas as pd, re
base=Path('outputs/category_debug')
map_path=base/'category_mapping.xlsx'
headers=['店铺','商品ID','商家编码','商品简称','一级品类','二级品类','是否启用','备注']
raw=pd.read_excel(map_path, dtype=str).fillna('')
# remove duplicate columns by position/name, then rebuild exact schema
raw=raw.loc[:, ~raw.columns.duplicated()].copy()
mapping=pd.DataFrame()
for h in headers:
    mapping[h]=raw[h].astype(str).str.strip() if h in raw.columns else ''
mapping['店铺']=mapping['店铺'].replace({'盲盒抖音':'盲盒抖店'})
pat=re.compile(r'(?i)(?:kuku|kukuna)|(?:^|[^0-9])(?:401|402|403|404|405|406)(?:[^0-9]|$)')
rows=[]
for store, filename in [('国货严选','guohuo.xlsx'), ('盲盒抖店','douyin.xlsx')]:
    wb=load_workbook(base/filename, data_only=True, read_only=False)
    for ws in wb.worksheets:
        data=[[cell.value for cell in row] for row in ws.iter_rows()]
        header_idx=None; idx={}
        for i,row in enumerate(data[:25]):
            vals=[str(v).strip() if v is not None else '' for v in row]
            if '商品ID' in vals and '货号' in vals and '单品结余' in vals:
                header_idx=i
                idx={name:j for j,name in enumerate(vals) if name}
                break
        if header_idx is None:
            continue
        pid_col=idx['商品ID']; sku_col=idx['货号']
        current_pid=''
        for row in data[header_idx+1:]:
            if pid_col < len(row) and str(row[pid_col] or '').strip():
                current_pid=str(row[pid_col]).strip()
            sku=str(row[sku_col] or '').strip() if sku_col < len(row) else ''
            if not current_pid or not sku or not pat.search(sku):
                continue
            rows.append([store,current_pid,sku,'','','','是',''])
    wb.close()
candidates=pd.DataFrame(rows, columns=headers).drop_duplicates(['店铺','商品ID','商家编码'])
print('candidate rows', len(candidates))
for store in ['国货严选','盲盒抖店']:
    sub=candidates[candidates['店铺'].eq(store)]
    print(store, len(sub), sorted(set(sum([re.findall(r'401|402|403|404|405|406', str(x)) for x in sub['商家编码']], []))))
known=set(zip(mapping['店铺'], mapping['商品ID'], mapping['商家编码'].str.upper().str.replace(' ', '', regex=False)))
add=[]
for _,r in candidates.iterrows():
    key=(r['店铺'], r['商品ID'], str(r['商家编码']).upper().replace(' ',''))
    if key not in known:
        add.append(r.tolist())
print('add rows', len(add))
merged=pd.concat([mapping, pd.DataFrame(add, columns=headers)], ignore_index=True) if add else mapping
out=base/'category_mapping_merged.xlsx'
merged.to_excel(out, index=False, sheet_name='品类映射')
print('wrote', out, len(merged))
for store in ['国货严选','盲盒抖店']:
    sub=merged[merged['店铺'].eq(store)]
    print('merged', store, len(sub), sorted(set(sum([re.findall(r'401|402|403|404|405|406', str(x)) for x in sub['商家编码']], []))))