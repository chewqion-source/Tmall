import pandas as pd, re
from pathlib import Path
from openpyxl import load_workbook
map_path=Path('D:/qianniu_profit_crawler/outputs/category_mapping_server.xlsx')
report_path=Path('D:/qianniu_profit_crawler/outputs/zy_report_server.xlsx')
print('mapping', map_path.exists(), map_path.stat().st_size)
m=pd.read_excel(map_path,dtype=str).fillna('')
print('columns', list(m.columns))
print('rows', len(m), 'stores', sorted(m['店铺'].astype(str).str.strip().unique().tolist()))
print('filled', int(m['一级品类'].astype(str).str.strip().ne('').sum()))
print(m.head(10).to_string(index=False))

def norm_store(s): return str(s or '').strip().replace('坐拥宁静','坐拥_宁静')
def norm_sku(s): return str(s or '').strip().upper().replace(' ','')
def report_date(sheet):
    mt=re.fullmatch(r'(\d{1,2})\.(\d{1,2})',str(sheet).strip())
    return pd.Timestamp(2026,int(mt.group(1)),int(mt.group(2))) if mt else None
def num(x):
    try: return float(str(x).replace(',','').replace('¥','').strip() or 0)
    except Exception: return 0.0
wb=load_workbook(report_path,data_only=True,read_only=False)
rows=[]
for ws in wb.worksheets:
    d=report_date(ws.title)
    if d is None: continue
    vals=[[c.value for c in row] for row in ws.iter_rows()]
    for mr in ws.merged_cells.ranges:
        min_col,min_row,max_col,max_row=mr.bounds; v=vals[min_row-1][min_col-1]
        for r in range(min_row-1,max_row):
            for c in range(min_col-1,max_col):
                if r<len(vals) and c<len(vals[r]): vals[r][c]=v
    cols=None; header=None
    for i,row in enumerate(vals[:25]):
        mp={str(v).strip():j for j,v in enumerate(row) if v is not None and str(v).strip()}
        if {'商品ID','数量','单品结余'}.issubset(mp): cols=mp; header=i; break
    if cols is None: continue
    product_col=cols['商品ID']; sku_col=cols.get('货号', cols.get('商家编码'))
    qty_col=cols['数量']; pay_col=cols.get('支付金额'); profit_col=cols['单品结余']
    refund_col=cols.get('成功退款金额') or cols.get('总退款金额')
    mx=max(x for x in [product_col,sku_col,qty_col,pay_col,profit_col,refund_col] if x is not None)
    for row in vals[header+1:]:
        if len(row)<=mx: continue
        pid='' if row[product_col] is None else str(row[product_col]).strip()
        sku='' if sku_col is None or row[sku_col] is None else str(row[sku_col]).strip()
        if not pid and not sku: continue
        if pid in {'合计','总计','小计'} or sku in {'合计','总计','小计'}: continue
        if num(row[qty_col])==0 and (pay_col is None or num(row[pay_col])==0) and num(row[profit_col])==0: continue
        rows.append({'date':d,'店铺':'坐拥_宁静','商品ID':pid,'商家编码':sku,'sku_norm':norm_sku(sku),'支付金额':num(row[pay_col]) if pay_col is not None else 0,'退款金额':num(row[refund_col]) if refund_col is not None else 0,'盈亏':num(row[profit_col])})
wb.close()
det=pd.DataFrame(rows)
print('detail rows', len(det), 'dates', det.date.min(), det.date.max())
m['店铺']=m['店铺'].map(norm_store); m['商品ID']=m['商品ID'].astype(str).str.strip(); m['sku_norm']=m['商家编码'].map(norm_sku)
m=m[(m['店铺']=='坐拥_宁静') & (m.get('是否启用','是').replace('', '是')!='否')]
joined=det.merge(m[['店铺','商品ID','sku_norm','商品简称','一级品类','二级品类']], on=['店铺','商品ID','sku_norm'], how='left')
joined['一级品类']=joined['一级品类'].fillna('').replace('', '未分类')
unique_total=joined[['商品ID','sku_norm']].drop_duplicates().shape[0]
unmapped=joined[joined['一级品类'].eq('未分类')][['商品ID','商家编码']].drop_duplicates()
print('unique sku', unique_total, 'unmapped', len(unmapped), 'matched', unique_total-len(unmapped))
print(joined.groupby('一级品类')[['支付金额','退款金额','盈亏']].sum().sort_values('支付金额',ascending=False).head(20).to_string())
print('unmapped sample')
print(unmapped.head(20).to_string(index=False))
